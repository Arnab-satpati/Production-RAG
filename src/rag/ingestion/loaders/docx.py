from __future__ import annotations

from docx import Document as DocxDocument

from rag.ingestion.loaders.base import BaseLoader, Document


class DocxLoader(BaseLoader):
    supported_extensions = [".docx"]

    def load(self) -> list[Document]:
        doc = DocxDocument(str(self.file_path))
        base_meta = self._base_metadata()

        current_heading = ""
        paragraphs_by_section: dict[str, list[str]] = {"": []}

        for para in doc.paragraphs:
            if para.style and para.style.name.startswith("Heading"):
                current_heading = para.text.strip()
                if current_heading not in paragraphs_by_section:
                    paragraphs_by_section[current_heading] = []
            elif para.text.strip():
                paragraphs_by_section.setdefault(current_heading, []).append(
                    para.text.strip(),
                )

        documents: list[Document] = []

        for heading, paragraphs in paragraphs_by_section.items():
            content = "\n\n".join(paragraphs)
            if content.strip():
                metadata = {
                    **base_meta,
                    "section_heading": heading,
                    "paragraph_count": len(paragraphs),
                }
                documents.append(Document(content=content, metadata=metadata))

        if not documents:
            full_text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
            if full_text:
                documents.append(Document(content=full_text, metadata=base_meta))

        return documents
